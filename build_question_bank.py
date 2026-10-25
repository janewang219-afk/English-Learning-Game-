from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "pep_grade3_first_term_english_question_bank.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, name="Arial", east_asia="Hiragino Sans GB", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, size=10.5, bold=False, color="000000", align=None, before=0, after=4, line=1.12):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_text(doc, text, style=None, size=10.5, bold=False, align=None, before=0, after=4, color="000000"):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    style_paragraph(p, size=size, bold=bold, color=color, align=align, before=before, after=after)
    return p


def add_mixed_paragraph(doc, pieces, size=10.5, after=4, before=0, left_indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(left_indent)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    for text, bold in pieces:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold)
    return p


def add_unit_header(doc, unit_no, title, question_focus, letters):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(f"Unit {unit_no} {title}")
    set_run_font(r, size=16, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    add_mixed_paragraph(
        doc,
        [("学习重点  ", True), (question_focus, False), ("  字母与语音  ", True), (letters, False)],
        size=10.5,
        after=8,
    )


def add_part_header(doc, label):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(label)
    set_run_font(r, size=12, bold=True)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True


def add_question(doc, number, text, answer_space=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.05)
    p.paragraph_format.first_line_indent = Inches(-0.05)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(f"{number}. {text}")
    set_run_font(r, size=10.5)
    if answer_space:
        p.paragraph_format.space_after = Pt(answer_space)
    return p


def add_answer_group(doc, label, answers):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(label)
    set_run_font(r, size=10.5, bold=True)
    r2 = p.add_run("  ")
    set_run_font(r2, size=10.5)
    r3 = p.add_run("；".join(f"{i}. {a}" for i, a in answers))
    set_run_font(r3, size=10.5)


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, "5B6571")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=9.5, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            if ridx % 2 == 1:
                set_cell_shading(cell, "F2F4F6")
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(value)
            set_run_font(r, size=9.5)
    return table


units = [
    {
        "no": 1,
        "title": "Making friends",
        "focus": "问候、介绍自己、认识朋友、身体部位和友好行为",
        "letters": "Aa-Dd",
        "parts": [
            ("Part A  Letters and sounds", [
                ("把字母按顺序补全  A  __  C  __。", "B，D"),
                ("选出首字母与 arm 相同的一项  A. apple  B. ear  C. dog。", "A"),
                ("选出首字母与 mouth 相同的一项  A. cat  B. map  C. dog。", "B"),
                ("从下列单词中找出以 D 开头的单词  ear  dog  hand。", "dog"),
            ]),
            ("Part B  Words", [
                ("“耳朵”用英语怎么说  A. eye  B. ear  C. arm。", "B"),
                ("We use our ______ to smile.  A. mouth  B. hand  C. ear。", "A"),
                ("I can ______ my toy with you.  A. share  B. goodbye  C. listen。", "A"),
                ("选出与 nice 同类的词  A. good  B. arm  C. toy。", "A"),
            ]),
            ("Part C  Sentences", [
                ("补全对话  —Hello! My name is Tom.  —Hi! ______ name is Lily.  A. My  B. I  C. Me。", "A"),
                ("—Nice to meet you.  —______  A. Goodbye.  B. Nice to meet you too.  C. Thank you.。", "B"),
                ("根据情境选择表达：第一次见到新同学时说什么  A. Nice to meet you.  B. Sit down.  C. Open your book.。", "A"),
                ("把句子补完整  I can listen, ______ and help.  A. smile  B. arm  C. mouth。", "A"),
                ("把句子补完整  We can share ______.  A. toys  B. eyes  C. ears。", "A"),
            ]),
            ("Part D  Reading", [
                ("阅读短文并判断正误。Hi, I am Amy. I meet Ben at school. We say hi and smile. Amy has a new toy. She shares it with Ben. Ben says thanks.  (1) Amy and Ben meet at school. (2) Amy shares a toy. (3) Ben says goodbye.。", "(1) T  (2) T  (3) F"),
                ("根据短文回答：What does Amy have?", "She has a new toy."),
                ("根据短文回答：Are Amy and Ben good friends?", "Yes, they are."),
            ]),
            ("Part E  Listening and speaking", [
                ("听老师或家长朗读后选择：Leo says his name is ______.  A. Leo  B. Ben  C. Tom。", "A"),
                ("听老师或家长朗读后选择：Leo can ______.  A. smile, listen and help  B. run and jump  C. draw and sing。", "A"),
                ("口语任务：用 3 句英语介绍自己，并向同伴问好。可使用 Hello! My name is … Nice to meet you.。", "参考：Hello! My name is Amy. Nice to meet you."),
                ("写一写：用英语写出“我的名字是 Amy。很高兴认识你。”", "参考：Hello! My name is Amy. Nice to meet you."),
            ]),
        ],
        "script": "听力脚本：Hello! My name is Leo. Nice to meet you. I have a toy. I can smile, listen and help. Goodbye!",
        "answers": [
            ("Part A", [(1, "B，D"), (2, "A"), (3, "B"), (4, "dog")]),
            ("Part B", [(5, "B"), (6, "A"), (7, "A"), (8, "A")]),
            ("Part C", [(9, "A"), (10, "B"), (11, "A"), (12, "A"), (13, "A")]),
            ("Part D", [(14, "(1) T  (2) T  (3) F"), (15, "She has a new toy."), (16, "Yes, they are.")]),
            ("Part E", [(17, "A"), (18, "A"), (19, "参考答案见题目"), (20, "参考：Hello! My name is Amy. Nice to meet you.")]),
        ],
    },
    {
        "no": 2,
        "title": "Different families",
        "focus": "家庭成员、介绍家人、this 和 that、理解不同家庭",
        "letters": "Ee-Hh",
        "parts": [
            ("Part A  Letters and sounds", [
                ("把字母按顺序补全  E  __  G  __。", "F，H"),
                ("选出首字母与 family 相同的一项  A. fish  B. apple  C. dog。", "A"),
                ("选出首字母与 grandma 相同的一项  A. hand  B. girl  C. ear。", "B"),
                ("从下列单词中找出以 H 开头的单词  aunt  hand  eye。", "hand"),
            ]),
            ("Part B  Words", [
                ("“奶奶”用英语怎么说  A. grandma  B. grandpa  C. aunt。", "A"),
                ("My mother’s son is my ______.  A. brother  B. sister  C. cousin。", "A"),
                ("“堂表兄弟姐妹”用英语怎么说  A. uncle  B. cousin  C. baby。", "B"),
                ("选出与 big 相反的词  A. small  B. same  C. nice。", "A"),
            ]),
            ("Part C  Sentences", [
                ("看见近处的奶奶，应该说  A. This is my grandma.  B. That is my grandma.  C. Is this my grandma?。", "A"),
                ("—Is that your brother?  —______  A. Yes, it is.  B. Yes, I do.  C. No, I don’t.。", "A"),
                ("—Is this your sister?  —No, ______ my cousin.  A. it’s  B. he’s  C. they’re。", "A"),
                ("补全句子  I ______ a small family.  A. have  B. am  C. is。", "A"),
                ("选择合适的句子完成意思：家庭大小不同，但家人的爱是相同的。  A. Families are different, but family love is the same.  B. Families are small and blue.  C. Family members can fly.。", "A"),
            ]),
            ("Part D  Reading", [
                ("阅读短文并判断正误。Leo lives with his dad and grandma. Mei lives with her mum, aunt and cousin. Their families are different. They all love their families.  (1) Leo lives with his grandma. (2) Mei lives with her uncle. (3) Both children love their families.。", "(1) T  (2) F  (3) T"),
                ("根据短文回答：Who lives with Mei?", "Her mum, aunt and cousin live with her."),
                ("根据短文回答：Are the two families the same?", "No, they are different."),
            ]),
            ("Part E  Listening and speaking", [
                ("听老师或家长朗读后选择：The speaker has a ______ family.  A. big  B. small  C. blue。", "A"),
                ("听老师或家长朗读后选择：The speaker lives with a mum, dad and ______.  A. brother  B. tiger  C. teacher。", "A"),
                ("口语任务：拿一张全家福或画一个简单家谱，用英语介绍至少 3 位家人。", "参考：This is my mum. This is my dad. This is my brother."),
                ("写一写：用 This is... 介绍你的妈妈、爸爸和兄弟或姐妹中的三位。", "参考：This is my mum. This is my dad. This is my brother."),
            ]),
        ],
        "script": "听力脚本：Look! This is my family. I have a big family. I live with my mum, dad and little brother. This is my grandma too. We love each other.",
        "answers": [
            ("Part A", [(1, "F，H"), (2, "A"), (3, "B"), (4, "hand")]),
            ("Part B", [(5, "A"), (6, "A"), (7, "B"), (8, "A")]),
            ("Part C", [(9, "A"), (10, "A"), (11, "A"), (12, "A"), (13, "A")]),
            ("Part D", [(14, "(1) T  (2) F  (3) T"), (15, "Her mum, aunt and cousin live with her."), (16, "No, they are different.")]),
            ("Part E", [(17, "A"), (18, "A"), (19, "参考答案见题目"), (20, "参考：This is my mum. This is my dad. This is my brother.")]),
        ],
    },
    {
        "no": 3,
        "title": "Amazing animals",
        "focus": "宠物、野生动物、动物特征、have 和 do 问答",
        "letters": "Ii-Ll",
        "parts": [
            ("Part A  Letters and sounds", [
                ("把字母按顺序补全  I  __  K  __。", "J，L"),
                ("选出首字母与 lion 相同的一项  A. leaf  B. panda  C. apple。", "A"),
                ("选出首字母与 rabbit 相同的一项  A. dog  B. red  C. tiger。", "B"),
                ("从下列单词中找出以 L 开头的单词  fish  lion  cat。", "lion"),
            ]),
            ("Part B  Words", [
                ("“宠物”用英语怎么说  A. pet  B. zoo  C. fox。", "A"),
                ("选出属于宠物的一项  A. cat  B. giraffe  C. elephant。", "A"),
                ("选出属于野生动物的一项  A. rabbit  B. panda  C. dog。", "B"),
                ("The giraffe is ______.  A. tall  B. small  C. blue。", "A"),
            ]),
            ("Part C  Sentences", [
                ("—Do you have a pet?  —______  A. Yes, I do.  B. Yes, it is.  C. Yes, I am.。", "A"),
                ("—Do you have a cat?  —No, ______.  A. I don’t  B. I’m not  C. it isn’t。", "A"),
                ("—What’s that?  —______  A. It’s a fox.  B. It’s fox.  C. They are fox.。", "A"),
                ("选择合适的回答：Let’s go to the zoo!  A. Great!  B. I’m five.  C. It’s red.。", "A"),
                ("补全句子  I ______ cats.  A. like  B. likes  C. liking。", "A"),
            ]),
            ("Part D  Reading", [
                ("阅读短文并判断正误。Tom has a pet cat called Coco. Coco is cute. Tom also likes dogs. On Sunday, Tom goes to the zoo with his friend. They see a panda, a monkey and a tall giraffe.  (1) Coco is a dog. (2) Tom likes dogs. (3) They see a giraffe.。", "(1) F  (2) T  (3) T"),
                ("根据短文回答：What pet does Tom have?", "He has a cat."),
                ("根据短文回答：What animals do they see at the zoo?", "A panda, a monkey and a giraffe."),
            ]),
            ("Part E  Listening and speaking", [
                ("听老师或家长朗读后选择：The speaker has a ______.  A. rabbit  B. panda  C. lion。", "A"),
                ("听老师或家长朗读后选择：The rabbit can ______.  A. hop  B. fly  C. swim。", "A"),
                ("口语任务：选择一种动物，用 2 至 3 句英语说出它的名字、特征或你的喜好。", "参考：This is a panda. It is cute. I like pandas."),
                ("写一写：用一句话描述一只可爱的熊猫。", "参考：This is a panda. It is cute. I like pandas."),
            ]),
        ],
        "script": "听力脚本：I have a pet rabbit. It is cute. It can hop. I like rabbits. At the zoo I see a fast lion.",
        "answers": [
            ("Part A", [(1, "J，L"), (2, "A"), (3, "B"), (4, "lion")]),
            ("Part B", [(5, "A"), (6, "A"), (7, "B"), (8, "A")]),
            ("Part C", [(9, "A"), (10, "A"), (11, "A"), (12, "A"), (13, "A")]),
            ("Part D", [(14, "(1) F  (2) T  (3) T"), (15, "He has a cat."), (16, "A panda, a monkey and a giraffe.")]),
            ("Part E", [(17, "A"), (18, "A"), (19, "参考答案见题目"), (20, "参考：This is a panda. It is cute. I like pandas.")]),
        ],
    },
    {
        "no": 4,
        "title": "Plants around us",
        "focus": "水果和植物、植物的需要、帮助植物生长、do 问答",
        "letters": "Mm-Pp",
        "parts": [
            ("Part A  Letters and sounds", [
                ("把字母按顺序补全  M  __  O  __。", "N，P"),
                ("选出首字母与 map 相同的一项  A. mum  B. orange  C. pen。", "A"),
                ("选出首字母与 orange 相同的一项  A. new  B. apple  C. pen。", "B"),
                ("从下列单词中找出以 P 开头的单词  fan  orange  pen。", "pen"),
            ]),
            ("Part B  Words", [
                ("“葡萄”用英语怎么说  A. grapes  B. grass  C. flowers。", "A"),
                ("Plants need air, ______ and sun.  A. water  B. toy  C. ear。", "A"),
                ("We can ______ the flowers.  A. water  B. listen  C. say。", "A"),
                ("Apple trees can give us ______.  A. apples  B. eyes  C. numbers。", "A"),
            ]),
            ("Part C  Sentences", [
                ("—Do you like apples?  —______  A. Yes, I do.  B. Yes, it is.  C. Yes, I am.。", "A"),
                ("—Do you like bananas?  —No, I ______.  A. don’t  B. am not  C. isn’t。", "A"),
                ("选择正确的句子  A. Plants need air, water and sun.  B. Plants needs air, water and sun.  C. Plant need airs.。", "A"),
                ("How can we help plants?  A. We can plant new trees.  B. We can eat the grass.  C. We can cut the flowers.。", "A"),
                ("根据中文补全  我们需要植物来生活。We need ______ to live.  A. plants  B. toys  C. colours。", "A"),
            ]),
            ("Part D  Reading", [
                ("阅读短文并判断正误。At our school garden, there is an apple tree and some flowers. Plants need air, water and sun. We water the flowers and plant new trees. The apple tree gives us apples.  (1) There is an apple tree in the garden. (2) Plants need only water. (3) The tree gives us apples.。", "(1) T  (2) F  (3) T"),
                ("根据短文回答：What do the plants need?", "Air, water and sun."),
                ("根据短文回答：What can the children do?", "They can water the flowers and plant new trees."),
            ]),
            ("Part E  Listening and speaking", [
                ("听老师或家长朗读后选择：The speaker likes ______.  A. grapes and oranges  B. bananas and grass  C. trees and toys。", "A"),
                ("听老师或家长朗读后选择：The speaker does not like ______.  A. bananas  B. apples  C. grapes。", "A"),
                ("口语任务：选择一种植物，说出它需要什么，以及我们能为它做什么。", "参考：Plants need air, water and sun. We can water the tree."),
                ("写一写：写出一句植物需要什么和一句我们能做什么。", "参考：Plants need air, water and sun. We can water the tree."),
            ]),
        ],
        "script": "听力脚本：I like grapes and oranges. I do not like bananas. There is a small tree in our garden. It needs air, water and sun.",
        "answers": [
            ("Part A", [(1, "N，P"), (2, "A"), (3, "B"), (4, "pen")]),
            ("Part B", [(5, "A"), (6, "A"), (7, "A"), (8, "A")]),
            ("Part C", [(9, "A"), (10, "A"), (11, "A"), (12, "A"), (13, "A")]),
            ("Part D", [(14, "(1) T  (2) F  (3) T"), (15, "Air, water and sun."), (16, "They can water the flowers and plant new trees.")]),
            ("Part E", [(17, "A"), (18, "A"), (19, "参考答案见题目"), (20, "参考：Plants need air, water and sun. We can water the tree.")]),
        ],
    },
    {
        "no": 5,
        "title": "The colourful world",
        "focus": "颜色、颜色混合、颜色的含义、询问颜色和喜好",
        "letters": "Qq-Uu",
        "parts": [
            ("Part A  Letters and sounds", [
                ("把字母按顺序补全  Q  __  S  __  U。", "R，T"),
                ("选出首字母与 quiet 相同的一项  A. queen  B. red  C. bus。", "A"),
                ("选出首字母与 red 相同的一项  A. ruler  B. see  C. up。", "A"),
                ("从下列单词中找出以 U 开头的单词  sit  up  run。", "up"),
            ]),
            ("Part B  Words", [
                ("“紫色”用英语怎么说  A. purple  B. pink  C. brown。", "A"),
                ("The sea is often ______.  A. blue  B. yellow  C. black。", "A"),
                ("A banana is often ______.  A. yellow  B. purple  C. brown。", "A"),
                ("选出表示“小心”的颜色信息  A. Yellow can say Be careful.  B. Blue can say No.  C. Red can say Go.。", "A"),
            ]),
            ("Part C  Sentences", [
                ("—What colour is it?  —______  A. It’s green.  B. They are green.  C. I like green.。", "A"),
                ("—What colours do you like?  —______  A. I like red and pink.  B. It’s red.  C. They are bears.。", "A"),
                ("Red and blue make ______.  A. purple  B. green  C. white。", "A"),
                ("Blue and yellow make ______.  A. green  B. orange  C. black。", "A"),
                ("根据短文提示选择：Green can say ______.  A. Go!  B. No!  C. Be careful!。", "A"),
            ]),
            ("Part D  Reading", [
                ("阅读短文并判断正误。At the park, the grass is green and the flowers are pink. A red sign tells us to stop. A green light tells us to go. A yellow sign tells us to be careful.  (1) The grass is green. (2) Red says Go. (3) Yellow tells us to be careful.。", "(1) T  (2) F  (3) T"),
                ("根据短文回答：What colour are the flowers?", "They are pink."),
                ("根据短文回答：What does a green light tell us?", "It tells us to go."),
            ]),
            ("Part E  Listening and speaking", [
                ("听老师或家长朗读后选择：The speaker sees a ______ bear.  A. brown  B. blue  C. pink。", "A"),
                ("听老师或家长朗读后选择：The speaker likes ______.  A. yellow and red  B. black and green  C. purple and brown。", "A"),
                ("口语任务：选择身边的 3 种物品，用 What colour is it? 和 It’s … 进行问答。", "参考：What colour is it? It’s blue."),
                ("写一写：用 What colour is it? 和 It’s... 写一组问答。", "参考：What colour is it? It’s blue."),
            ]),
        ],
        "script": "听力脚本：I see a brown bear, a blue sea and a purple flower. I like yellow and red. Colours are beautiful and useful.",
        "answers": [
            ("Part A", [(1, "R，T"), (2, "A"), (3, "A"), (4, "up")]),
            ("Part B", [(5, "A"), (6, "A"), (7, "A"), (8, "A")]),
            ("Part C", [(9, "A"), (10, "A"), (11, "A"), (12, "A"), (13, "A")]),
            ("Part D", [(14, "(1) T  (2) F  (3) T"), (15, "They are pink."), (16, "It tells us to go.")]),
            ("Part E", [(17, "A"), (18, "A"), (19, "参考答案见题目"), (20, "参考：What colour is it? It’s blue.")]),
        ],
    },
    {
        "no": 6,
        "title": "Useful numbers",
        "focus": "数字 1 至 10、年龄、数量、整点、钱和生日",
        "letters": "Vv-Zz",
        "parts": [
            ("Part A  Letters and sounds", [
                ("把字母按顺序补全  V  __  X  __  Z。", "W，Y"),
                ("选出首字母与 van 相同的一项  A. vet  B. six  C. yellow。", "A"),
                ("选出首字母与 yellow 相同的一项  A. yo-yo  B. box  C. quiz。", "A"),
                ("选出单词末尾音与 box 相同的一项  A. six  B. van  C. vet。", "A"),
            ]),
            ("Part B  Words", [
                ("数字 7 的英语是  A. six  B. seven  C. eight。", "B"),
                ("数字 9 的英语是  A. nine  B. ten  C. five。", "A"),
                ("“几岁”用英语怎么问  A. How old are you?  B. How many are you?  C. What colour are you?。", "A"),
                ("整点 7 点可以说  A. It’s seven o’clock.  B. It’s seven years.  C. Seven yuan, please.。", "A"),
            ]),
            ("Part C  Sentences", [
                ("—How old are you?  —______  A. I’m eight years old.  B. I have eight.  C. It’s eight o’clock.。", "A"),
                ("—How many apples?  —______  A. Three.  B. I’m three.  C. It’s red.。", "A"),
                ("购物时说“十元，谢谢”可以说  A. That’s ten yuan, please.  B. Ten years old, please.  C. Ten apples, please.。", "A"),
                ("别人递给你东西时，你可以说  A. Here you are.  B. What colour is it?  C. How old are you?。", "A"),
                ("生日聚会邀请中，最合适的表达是  A. Please come to my birthday party.  B. Please water my tree.  C. Please go to the zoo.。", "A"),
            ]),
            ("Part D  Reading", [
                ("阅读短文并判断正误。It is Lili’s birthday. She is seven years old. At seven o’clock, her friends come. There are one cake, two apples and three bananas.  (1) Lili is seven. (2) Her friends come at eight o’clock. (3) There are three bananas.。", "(1) T  (2) F  (3) T"),
                ("根据短文回答：How many apples are there?", "There are two apples."),
                ("根据短文回答：What is the event?", "It is Lili’s birthday."),
            ]),
            ("Part E  Listening and speaking", [
                ("听老师或家长朗读后选择：The speaker is ______ years old.  A. six  B. four  C. ten。", "A"),
                ("听老师或家长朗读后选择：The speaker sees ______ birds.  A. eight  B. nine  C. two。", "A"),
                ("口语任务：制作一个简单生日邀请卡，至少说出姓名、年龄和聚会时间中的两项。", "参考：I’m six years old. Please come to my birthday party."),
                ("写一写：用英语说出年龄并邀请朋友参加生日聚会。", "参考：I’m six years old. Please come to my birthday party."),
            ]),
        ],
        "script": "听力脚本：I am six years old. I have four pencils. I see eight birds in the tree. It is five o’clock now.",
        "answers": [
            ("Part A", [(1, "W，Y"), (2, "A"), (3, "A"), (4, "A")]),
            ("Part B", [(5, "B"), (6, "A"), (7, "A"), (8, "A")]),
            ("Part C", [(9, "A"), (10, "A"), (11, "A"), (12, "A"), (13, "A")]),
            ("Part D", [(14, "(1) T  (2) F  (3) T"), (15, "There are two apples."), (16, "It is Lili’s birthday.")]),
            ("Part E", [(17, "A"), (18, "A"), (19, "参考答案见题目"), (20, "参考：I’m six years old. Please come to my birthday party.")]),
        ],
    },
]


final_questions = [
    ("字母  A, C, E, __, I。", "G"),
    ("字母  M, N, __, P。", "O"),
    ("选出首字母与 rabbit 相同的单词  A. red  B. dog  C. apple。", "A"),
    ("选出单词末尾音与 box 相同的单词  A. six  B. vet  C. yellow。", "A"),
    ("把单词按类别归类：grandma, panda, grapes, blue。Family member: ______；Animal: ______；Fruit: ______；Colour: ______。", "grandma；panda；grapes；blue"),
    ("“帮助”用英语怎么说  A. help  B. share  C. have。", "A"),
    ("“表兄弟姐妹”用英语怎么说  A. cousin  B. uncle  C. brother。", "A"),
    ("植物需要 air, ______ and sun。", "water"),
    ("Red and blue make ______。", "purple"),
    ("数字 10 的英语是 ______。", "ten"),
    ("—Nice to meet you.  —______。", "Nice to meet you too."),
    ("—Is that your sister?  —No, ______ my cousin。", "it’s"),
    ("—Do you have a pet?  —Yes, ______。", "I do"),
    ("—What colour is it?  —______ green。", "It’s"),
    ("—How old are you?  —I’m ______ years old。", "答案可为 five / six / seven / eight 等合理年龄"),
    ("根据情境补全：We can ______ the flowers.  A. water  B. eat  C. hear。", "A"),
    ("根据情境补全：This is ______ grandma.  A. my  B. I  C. me。", "A"),
    ("根据情境补全：I ______ cats.  A. like  B. likes  C. liking。", "A"),
    ("根据情境补全：How ______ apples?  A. many  B. old  C. colour。", "A"),
    ("根据情境补全：Please come to my birthday ______.  A. party  B. garden  C. zoo。", "A"),
    ("阅读：Mina has a small family. She lives with her mum and brother. She has a pet cat. The cat is black and cute. Mina likes her family and her cat. 判断：Mina has a big family. / The cat is cute. / Mina likes her cat。", "F；T；T"),
    ("根据上文回答：Who does Mina live with?", "Her mum and brother."),
    ("根据上文回答：What colour is the cat?", "It is black."),
    ("根据上文回答：What pet does Mina have?", "She has a cat."),
    ("根据上文给短文取一个合适的小标题。", "参考：Mina and her family / Mina’s cat"),
    ("听力：听朗读后选择。The child has a ______ family.  A. big  B. small  C. new。", "A"),
    ("听力：听朗读后选择。The child likes ______.  A. apples  B. purple  C. lions。", "A"),
    ("听力：听朗读后选择。The child’s pet can ______.  A. hop  B. fly  C. read。", "A"),
    ("听力：听朗读后选择。The birthday party is at ______ o’clock.  A. six  B. seven  C. ten。", "B"),
    ("综合口语：完成一段 5 句左右的小对话，至少包含一次问候、一次家庭或动物介绍、一次颜色或数字表达。", "参考：Hi! My name is Amy. Nice to meet you. This is my brother. He has a cute cat. It is black and he is six years old."),
]


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size, before, after in [
        ("Heading 1", 16, 10, 5),
        ("Heading 2", 12, 6, 3),
        ("Heading 3", 11, 4, 2),
    ]:
        st = styles[style_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("PEP 三年级上册英语全年题库  |  ")
    set_run_font(r, size=8.5, color="666666")
    add_page_number(p)
    for run in p.runs:
        set_run_font(run, size=8.5, color="666666")


def add_answer_page(doc):
    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("家长和老师用答案")
    set_run_font(r, size=16, bold=True)
    p.paragraph_format.space_after = Pt(5)
    add_text(doc, "答案中的开放题只给出示例，孩子表达正确且符合题意即可。听力题请先遮住本页脚本，由家长或老师朗读两遍。", size=10.5, after=8)
    for unit in units:
        p = doc.add_paragraph(style="Heading 2")
        r = p.add_run(f"Unit {unit['no']} {unit['title']}")
        set_run_font(r, size=12, bold=True)
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        for label, answers in unit["answers"]:
            add_answer_group(doc, label, answers)
        add_text(doc, unit["script"], size=9.5, color="555555", after=5)
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run("全册综合复习")
    set_run_font(r, size=12, bold=True)
    p.paragraph_format.space_before = Pt(8)
    add_answer_group(doc, "答案", [(i, a) for i, (_, a) in enumerate(final_questions, 1)])
    add_text(doc, "综合听力脚本：This is Ben. He has a big family. He likes apples. He has a pet rabbit. The rabbit can hop. Ben is seven years old. His birthday party is at seven o’clock.", size=9.5, color="555555", after=5)
    add_text(doc, "口语评分建议：每项 1 分，共 4 分。问候自然 1 分；句子基本完整 1 分；词汇或句型使用正确 1 分；发音和表达清楚 1 分。", size=9.5, color="555555", after=5)


def build():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    add_footer(section)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("三年级上册英语全年题库")
    set_run_font(r, size=24, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(18)
    r = sub.add_run("人教 PEP 2024 新版  个人学习练习用")
    set_run_font(r, size=12, color="555555")

    add_text(doc, "这套题库主要根据教材末尾的“Words in each unit”“Vocabulary”和“Useful expressions”三部分整理，覆盖六个学习单元和期末综合复习。题目把末尾词汇、全册字母、常用句型和情境表达放进原创的选择、阅读、听力和口语练习中，适合孩子在一学年中分单元使用。", size=11, after=8)
    add_text(doc, "建议用法  每单元先完成练习册部分，再让家长或老师朗读听力脚本；口语题不要求唯一答案，重点是敢说、说完整、意思对。", size=10.5, after=10)

    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("全册覆盖地图")
    set_run_font(r, size=16, bold=True)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    rows = [
        ("Unit 1", "Making friends", "问候、介绍、身体部位、友好行为", "Aa-Dd"),
        ("Unit 2", "Different families", "家庭成员、this 和 that、家庭表达", "Ee-Hh"),
        ("Unit 3", "Amazing animals", "宠物、野生动物、have 和 do 问答", "Ii-Ll"),
        ("Unit 4", "Plants around us", "水果、植物需求、帮助植物", "Mm-Pp"),
        ("Unit 5", "The colourful world", "颜色、混色、颜色的含义", "Qq-Uu"),
        ("Unit 6", "Useful numbers", "数字、年龄、数量、时间、生日", "Vv-Zz"),
        ("Revision", "Being a good guest", "综合运用和礼貌表达", "A-Z"),
    ]
    make_table(doc, ["单元", "主题", "重点内容", "字母"], rows, widths=[0.75, 1.65, 3.45, 0.75])
    add_text(doc, "题量说明  六个单元各 20 题，全册综合复习 30 题，共 150 题。建议每题 1 分；阅读、听力和口语可按孩子的表达情况灵活评分。", size=10, color="555555", after=10)
    add_text(doc, "教材入口  https://book.pep.com.cn/1212001301245/mobile/index.html", size=9.5, color="555555", after=4)

    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("练习册部分")
    set_run_font(r, size=16, bold=True)
    p.paragraph_format.space_after = Pt(5)
    add_text(doc, "请先独立完成。听力题的录音脚本放在后面的答案部分，口语题请尽量开口完成。", size=10.5, color="555555", after=8)

    for unit in units:
        add_unit_header(doc, unit["no"], unit["title"], unit["focus"], unit["letters"])
        qn = 1
        for label, questions in unit["parts"]:
            add_part_header(doc, label)
            for text, _ in questions:
                add_question(doc, qn, text)
                qn += 1

    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run("全册综合复习")
    set_run_font(r, size=16, bold=True)
    p.paragraph_format.space_after = Pt(5)
    add_text(doc, "完成六个单元后使用。综合题会交叉考查字母、词汇、句型和生活场景。", size=10.5, color="555555", after=8)
    add_part_header(doc, "Final review  30 questions")
    for i, (text, _) in enumerate(final_questions, 1):
        add_question(doc, i, text)

    add_answer_page(doc)

    # Keep all heading and table text black or explicitly styled.
    doc.core_properties.title = "三年级上册英语全年题库"
    doc.core_properties.subject = "人教 PEP 2024 新版英语三年级上册练习题"
    doc.core_properties.author = "OpenAI"
    doc.save(OUT)


if __name__ == "__main__":
    build()
